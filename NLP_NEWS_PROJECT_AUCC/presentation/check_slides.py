import sys, re
from pathlib import Path
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'tools'))
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation

files=sorted((ROOT/'preview').glob('*.PNG'), key=lambda p:int(re.search(r'(\d+)',p.stem).group(1)))
if not files:
    files=sorted((ROOT/'preview').glob('*.png'), key=lambda p:int(re.search(r'(\d+)',p.stem).group(1)))
sheet=Image.new('RGB',(1920,1160),'#dfe5ef')
for i,f in enumerate(files):
    im=Image.open(f).convert('RGB'); im.thumbnail((470,265))
    x=(i%4)*480+5; y=(i//4)*290+18
    sheet.paste(im,(x,y)); ImageDraw.Draw(sheet).text((x+4,y-15),str(i+1),fill='#14213d')
sheet.save(ROOT/'slide-overview.jpg',quality=92)
prs=Presentation(ROOT/'AUCC_Sports_Comment_Analysis.pptx')
issues=[]
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+1000 or sh.top+sh.height>prs.slide_height+1000:
            issues.append((i,sh.name,'out of bounds'))
    assert s.has_notes_slide, i
print('Slides:',len(prs.slides),'Preview images:',len(files),'Out-of-bounds:',issues)
print('All slides include speaker notes.')
