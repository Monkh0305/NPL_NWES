"""Build an editable Thai undergraduate project presentation from verified evidence."""
import sys, json, re
from pathlib import Path
ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT/'tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

E = json.loads((ROOT/'evidence.json').read_text(encoding='utf-8'))
R = E['result']
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'ระบบวิเคราะห์ความคิดเห็นต่อวิดีโอกีฬาบน YouTube ด้วย Gemini'
prs.core_properties.subject = 'โครงงานระดับปริญญาตรีสำหรับเตรียมนำเสนอ AUCC'
prs.core_properties.author = 'ผู้จัดทำโครงงาน'
prs.core_properties.keywords = 'YouTube, Gemini, Sentiment Analysis, Sports, AUCC'
FONT = 'Leelawadee UI'
NAVY='14213D'; DARK='0D172C'; INK='18243C'; MUTED='62718B'; BG='F4F6FB'
PURPLE='7155E8'; TEAL='009D83'; RED='E34B69'; LINE='DCE3EF'; WHITE='FFFFFF'; LIGHT='EDE8FF'
notes=[]

def color(h): return RGBColor.from_string(h)
def box(s,x,y,w,h,fill=WHITE,line=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill)
    sh.line.fill.background() if line is None else None
    if line: sh.line.color.rgb=color(line)
    if radius: sh.adjustments[0]=0.12
    return sh
def text(s,txt,x,y,w,h,size=22,bold=False,c=INK,align=None):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=Pt(2); tf.margin_bottom=0
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color(c)
        p.space_after=Pt(5)
        if align is not None: p.alignment=align
    return sh
def label(s,txt,x,y,w=2.4,fill=LIGHT,c=PURPLE):
    box(s,x,y,w,.35,fill,radius=True); text(s,txt,x+.1,y+.015,w-.2,.32,11,True,c)
def line(s,x1,y1,x2,y2,c=LINE,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); sh.line.color.rgb=color(c); sh.line.width=Pt(width)
def arrow(s,x,y,w=.4,h=.3,c=PURPLE): boxshape=s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)); boxshape.fill.solid(); boxshape.fill.fore_color.rgb=color(c); boxshape.line.fill.background()
def slide(title,section,sub=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=color(DARK if dark else BG)
    box(s,.5,.4,.08,.35,TEAL)
    text(s,section.upper(),.75,.4,10,.3,11,True,'A7B5D1' if dark else MUTED)
    text(s,title,.65,.94,12,1.0,30,True,WHITE if dark else INK)
    if sub: text(s,sub,.68,1.91,12,.62,16,False,'BCC8DE' if dark else MUTED)
    line(s,.65,7.02,12.68,7.02,'34415C' if dark else LINE,1)
    text(s,'YOUTUBE SPORTS COMMENT ANALYSIS  /  AUCC',.68,7.11,11,.2,9,False,'A7B5D1' if dark else MUTED)
    text(s,f'{len(prs.slides):02d}',12.1,7.08,.55,.27,12,True,WHITE if dark else PURPLE,PP_ALIGN.RIGHT)
    return s
def note(s,txt):
    s.notes_slide.notes_text_frame.text=txt
    notes.append((len(prs.slides),txt))
def card(s,x,y,w,h,title,body,num=None,c=PURPLE):
    box(s,x,y,w,h,WHITE,radius=True)
    box(s,x,y,.055,h,c)
    offset=.3
    if num:
        text(s,num,x+.23,y+.15,w-.45,.6,30,True,c); offset=.94
    text(s,title,x+.25,y+offset,w-.5,.75,22,True)
    text(s,body,x+.25,y+offset+.84,w-.5,h-offset-1,18,False,MUTED)
def photo(s,name,x,y,w,h,crop=None):
    iw,ih=Image.open(ROOT/'assets'/name).size
    if crop:
        iw*=1-crop[0]-crop[2]; ih*=1-crop[1]-crop[3]
    ratio=iw/ih
    fitted_w=min(w,h*ratio); fitted_h=fitted_w/ratio
    x+=(w-fitted_w)/2; y+=(h-fitted_h)/2; w,h=fitted_w,fitted_h
    sh=s.shapes.add_picture(str(ROOT/'assets'/name),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if crop: sh.crop_left,sh.crop_top,sh.crop_right,sh.crop_bottom=crop
    return sh

# 01 — cover
s=slide('','UNDERGRADUATE PROJECT • ORAL PRESENTATION',dark=True)
label(s,'NLP × SPORTS × WEB APPLICATION',.7,1.13,4.3,fill='253451',c='C7BAFF')
text(s,'ระบบวิเคราะห์ความคิดเห็น\nต่อวิดีโอกีฬาบน YouTube\nด้วย Gemini',.7,1.88,7.2,2.48,34,True,WHITE)
text(s,'YouTube Sports Comment Analysis\nwith Gemini',.74,4.6,7,.8,20,False,'B6C5DF')
text(s,'ผู้จัดทำ: [ชื่อ–นามสกุล / รหัสนักศึกษา]\nอาจารย์ที่ปรึกษา: [ชื่ออาจารย์]\n[สาขาวิชา / มหาวิทยาลัย]',.74,5.7,7,1.08,15,False,'C9D3E5')
box(s,8.55,1.9,4.05,4.7,'1C2A47',radius=True)
text(s,'COMMENTS',8.87,2.2,3.4,.4,13,True,'B6C5DF')
for i,(st,c) in enumerate([('ชื่นชมผลงาน',TEAL),('ตั้งคำถาม',PURPLE),('สะท้อนปัญหา',RED)]):
    box(s,8.9,2.89+i*.65,3.32,.48,'293959',radius=True); box(s,9.05,3.02+i*.65,.13,.13,c); text(s,st,9.38,2.94+i*.65,2.55,.38,17,False,WHITE)
arrow(s,10.26,4.99,.4,.34,TEAL)
text(s,'ข้อมูล → ภาพรวม',8.97,5.69,3.3,.5,25,True,WHITE)
note(s,'เวลาแนะนำ 30 วินาที แนะนำชื่อโครงงาน ผู้จัดทำ และอาจารย์ที่ปรึกษา งานนี้ประยุกต์ Gemini ผ่าน API เพื่อช่วยอ่านภาพรวมความคิดเห็นต่อวิดีโอกีฬา ไม่ใช่การฝึกโมเดลใหม่ เติมข้อมูลผู้จัดทำก่อนส่ง และตรวจรูปแบบกับงาน AUCC รุ่นที่จะส่งอีกครั้ง')

# 02 — problem
s=slide('จากความคิดเห็นรายข้อความ สู่ภาพรวมที่อ่านได้','01 / PROBLEM','โจทย์ของโครงงาน: จะรวบรวมและสรุปความคิดเห็นต่อวิดีโอกีฬาอย่างเป็นระบบได้อย่างไร')
card(s,.7,2.75,3.85,3.5,'ความคิดเห็นหลากหลาย','มีทั้งคำชม คำถาม และข้อวิจารณ์\nการอ่านทีละข้อความทำให้\nเปรียบเทียบภาพรวมได้ยาก','01')
card(s,4.75,2.75,3.85,3.5,'บริบทเฉพาะของกีฬา','กล่าวถึงนักกีฬา ทีม กรรมการ\nและผลการแข่งขัน\nหนึ่งข้อความมีหลายประเด็น','02',TEAL)
card(s,8.8,2.75,3.85,3.5,'ต้องสรุปอย่างมีขอบเขต','ต้องบอกจำนวนตัวอย่าง\nและแยกความเห็นของผู้ชม\nออกจากข้อเท็จจริง','03',RED)
note(s,'เวลาแนะนำ 45 วินาที อธิบายแรงจูงใจของการพัฒนาระบบ ไม่อ้างตัวเลขปริมาณข้อมูลหรือเวลาที่ประหยัดได้โดยยังไม่ได้ทดลอง เน้นความต่างระหว่างการอ่านแต่ละข้อความกับการเห็นสัดส่วนรวม')

# 03 — objectives
s=slide('วัตถุประสงค์และขอบเขตของระบบ','02 / OBJECTIVES & SCOPE')
for y,n,t in [(2.27,'1','ค้นหาและเลือกวิดีโอกีฬาจาก YouTube'),(3.38,'2','จำแนกความรู้สึก อารมณ์ และประเด็นที่กล่าวถึง'),(4.49,'3','แสดงกราฟ สรุปภาษาไทย และจัดเก็บผลวิเคราะห์')]:
    label(s,n,.76,y,.48,fill=PURPLE,c=WHITE); text(s,t,1.45,y-.015,7.1,.84,23,True)
box(s,9.0,2.27,3.55,4.27,NAVY,radius=True)
text(s,'ขอบเขตข้อมูล',9.3,2.58,2.95,.45,22,True,WHITE)
text(s,'≤ 100',9.3,3.15,2.95,.72,40,True,'9CE8D9')
text(s,'ความคิดเห็นหลักต่อรอบ',9.3,3.9,2.95,.54,17,False,WHITE)
text(s,'≤ 2,000 ตัวอักษร/ข้อความ\nภาษาไทยและอังกฤษ\nหมวด Sports (รหัส 17)',9.3,4.62,2.95,1.55,17,False,'C6D2E8')
text(s,'ไม่รวมการดึงทุก reply และไม่อ้างว่าเป็นตัวแทนผู้ชมทั้งหมด',.78,6.26,7.9,.45,16,False,MUTED)
note(s,'เวลาแนะนำ 45 วินาที ใช้ top-level comments ตามที่ API ส่งมา ไม่ใช่การสุ่มตัวอย่างเชิงสถิติ และไม่ใช่ทุกความคิดเห็นรวม replies ขนาดจริงอาจน้อยกว่า 100 ถ้ามีข้อมูลน้อยกว่า อ้างอิง app/config.py และ youtube_service.py')

# 04 — architecture
s=slide('สถาปัตยกรรม: แยกหน้าจอ การวิเคราะห์ และแหล่งข้อมูล','03 / SYSTEM ARCHITECTURE')
for x,y,w,h,title,body,fill in [(.72,3.18,2.65,1.6,'React + Vite','แดชบอร์ด / กราฟ',WHITE),(4.1,3.18,3.0,1.6,'FastAPI','ตรวจผล / รวมสถิติ',NAVY),(8.5,2.32,3.93,1.12,'YouTube Data API','ข้อมูลวิดีโอและความคิดเห็น',WHITE),(8.5,3.75,3.93,1.12,'Gemini API','Structured JSON',WHITE),(4.1,5.47,3,1.0,'SQLite','บันทึกผลแต่ละครั้ง',WHITE)]:
    box(s,x,y,w,h,fill,radius=True); text(s,title,x+.22,y+.19,w-.4,.48,23,True,WHITE if fill==NAVY else INK); text(s,body,x+.22,y+.76,w-.4,.43,15,False,'C6D2E8' if fill==NAVY else MUTED)
arrow(s,3.53,3.77,.38,.32); line(s,7.1,3.98,7.79,3.98,PURPLE); line(s,7.79,2.86,7.79,4.27,PURPLE); line(s,7.79,2.86,8.42,2.86,PURPLE); line(s,7.79,4.27,8.42,4.27,PURPLE); line(s,5.6,4.79,5.6,5.45,TEAL)
label(s,'HTTP / JSON',1.24,5.13,1.75)
text(s,'API keys อยู่ฝั่ง backend',8.52,5.46,3.93,.55,20,True,TEAL)
text(s,'Cloudflare Tunnel ใช้เปิดทางเข้าจากอินเทอร์เน็ต → frontend → /api → backend',.8,6.64,11.8,.27,13,False,MUTED)
note(s,'เวลาแนะนำ 60 วินาที Frontend ใช้ React, Tailwind และ Recharts ติดต่อ FastAPI ผ่าน /api Backend ติดต่อ YouTube และ Gemini โดยไม่ส่ง API key ไปหน้าเว็บ ผลถูกเก็บด้วย SQLAlchemy ใน SQLite ไม่มีการประมวลผลโมเดลภายในเครื่อง Cloudflare เป็นทางเชื่อมเข้าเครื่องที่รันระบบ ไม่ใช่ที่รันโมเดล')

# 05 — workflow
s=slide('กระบวนการวิเคราะห์ตั้งแต่เลือกวิดีโอจนแสดงผล','04 / ANALYSIS PIPELINE')
steps=[('เลือกวิดีโอ','ค้นหา / เลือกหมวด\nจากรายการกีฬา'),('ดึงความคิดเห็น','เก็บข้อความหลัก\nสูงสุด 100 รายการ'),('เรียก Gemini','กำหนด prompt\nและ JSON schema'),('ตรวจและรวมผล','ตรวจ ID / ประเภท\nคำนวณเปอร์เซ็นต์'),('บันทึกและแสดง','เก็บ snapshot\nแสดงกราฟ + สรุป')]
for i,(t,b) in enumerate(steps):
    x=.7+i*2.52; box(s,x,2.93,2.31,2.75,WHITE,radius=True); label(s,f'0{i+1}',x+.18,3.12,.55,fill=PURPLE,c=WHITE); text(s,t,x+.17,3.81,1.98,.5,20,True); text(s,b,x+.17,4.51,1.98,.87,16,False,MUTED)
    if i<4: arrow(s,x+2.35,4.05,.14,.23)
box(s,.73,6.04,11.86,.57,LIGHT,radius=True); text(s,'ข้อมูลครบและรูปแบบถูกต้องก่อนบันทึกผล — หากผิดพลาด แจ้งเหตุผลให้ผู้ใช้',.98,6.11,11.36,.38,19,True,PURPLE)
note(s,'เวลาแนะนำ 50 วินาที ลำดับนี้อ้างจาก _run_analysis ใน routers/analysis.py หลังเรียก Gemini จะตรวจผลครบก่อนคำนวณสถิติและ commit ฐานข้อมูล หน้าจอโหลดรายละเอียดวิดีโอขนานกับการวิเคราะห์ แต่รอความคิดเห็นยอดนิยมอีกครั้งก่อนแสดงผลทั้งหมด')

# 06 — output and validation
s=slide('ผลลัพธ์มีโครงสร้าง และตรวจสอบก่อนคำนวณ','05 / STRUCTURED OUTPUT')
for x,n,t,b,c in [(.74,'3','Sentiment','บวก / เป็นกลาง / ลบ',TEAL),(4.9,'6','Emotion','อารมณ์หลัก 1 กลุ่มต่อข้อความ',PURPLE),(9.06,'12','Topics','หลายประเด็นต่อข้อความได้',RED)]:
    box(s,x,2.5,3.52,1.95,WHITE,radius=True); text(s,n,x+.2,2.69,.95,.85,44,True,c); text(s,t,x+1.22,2.76,2.03,.42,21,True); text(s,b,x+.22,3.65,3.06,.6,15,False,MUTED)
text(s,'ตรวจด้วย JSON Schema + Pydantic',.82,4.93,5.95,.45,22,True)
text(s,'ID ของความคิดเห็นต้องครบและไม่ซ้ำ\nค่า label ต้องอยู่ในชุดที่กำหนด\nผลผิดรูปแบบ: ลองใหม่ได้อีก 1 ครั้ง',.82,5.53,5.7,1.15,18,False,MUTED)
box(s,7.2,4.88,5.35,1.83,NAVY,radius=True)
text(s,'สัดส่วน = (จำนวนในกลุ่ม ÷ n) × 100',7.46,5.12,4.88,.67,20,True,WHITE)
text(s,'Topics รวมเกิน 100% ได้ เพราะเป็นหลายป้ายกำกับ',7.46,6.0,4.86,.5,15,False,'C6D2E8')
note(s,'เวลาแนะนำ 60 วินาที อธิบาย sentiment เทียบกับ emotion: sentiment เป็นแนวโน้มบวกกลางลบ ส่วน emotion เป็นยินดี โกรธ เศร้า ประหลาดใจ กลัว และเป็นกลาง Topics มี 12 กลุ่มแต่จำนวนที่พบในแต่ละวิดีโออาจน้อยกว่า ทุกเปอร์เซ็นต์นับจากรายการที่ตรวจครบแล้ว ไม่ใช่ confidence ของโมเดล Structured output ยืนยันโครงสร้าง ไม่รับประกันความถูกต้องทางความหมาย')

# 07 — dashboard screenshot
s=slide('หน้าค้นหาและเลือกวิดีโอ','06 / SYSTEM DEMONSTRATION','ภาพหน้าจอจากระบบจริง • ข้อมูลจาก YouTube Data API')
photo(s,'dashboard.png',.72,2.62,8.48,4.16,crop=(.055,.085,.055,.12))
for y,n,t,b in [(2.73,'01','เลือกหมวดกีฬา','หรือพิมพ์คำค้น'),(4.05,'02','เลือกรายการ','ยอดนิยม / ล่าสุด / ยอดชม'),(5.37,'03','กดวิเคราะห์','เปิดหน้าผลของวิดีโอนั้น')]:
    label(s,n,9.52,y,.55); text(s,t,10.26,y-.01,2.24,.44,21,True); text(s,b,9.56,y+.57,2.96,.55,16,False,MUTED)
note(s,'เวลาแนะนำ 45 วินาที อธิบายองค์ประกอบหน้าจอและเส้นทางใช้งาน ภาพนี้จับจากระบบจริงในวันที่ 27 กันยายน 2569 รายการยอดนิยมและอันดับแสดงตามข้อมูลที่ระบบดึงมา ไม่ใช่อันดับกีฬาทั้งโลก ภาพปกและชื่อวิดีโอเป็นของเจ้าของเนื้อหาบน YouTube ใช้ประกอบการสาธิตเชิงการศึกษา')

# 08 — analysis screenshot
s=slide('หน้าวิเคราะห์: อ่านสัดส่วนความรู้สึกและอารมณ์','07 / ANALYSIS DASHBOARD','ภาพจริงของ UI โดยแสดงผลวิเคราะห์ที่บันทึกไว้ • 27 ความคิดเห็นหลัก')
photo(s,'analysis.png',.72,2.63,9.1,4.07,crop=(.06,.407,.061,.13))
text(s,'อ่านผลอย่างไร',10.1,2.83,2.5,.6,22,True)
text(s,'วงแหวน\nสัดส่วน sentiment\n\nแถบแนวนอน\nสัดส่วนอารมณ์หลัก\n\nจำนวนตัวอย่าง\nบอกฐานของเปอร์เซ็นต์',10.1,3.6,2.5,2.85,17,False,MUTED)
note(s,'เวลาแนะนำ 50 วินาที ภาพนี้ใช้ผลจริงจาก AnalysisSnapshot ที่บันทึกไว้วันที่ 27 กันยายน 2569 เวลา 12:03:19 UTC โดย replay response เฉพาะระหว่างจับภาพ ไม่เรียก inference ใหม่ กราฟในภาพจึงเป็นผลที่มีหลักฐาน ไม่ใช่ภาพจำลองหรือข้อมูลสมมติ อธิบาย n=27 และแยก sentiment กับ emotion')

# 09 — case study chart
s=slide('กรณีสาธิต: วิเคราะห์ความคิดเห็น 27 รายการ','08 / OBSERVED OUTPUT','วิดีโอสรุปการแข่งขันเอเชียนเกมส์ • ช่อง T Sports 7 • ผลที่บันทึก 27 กันยายน 2569')
data=CategoryChartData(); data.categories=['เชิงบวก','เป็นกลาง','เชิงลบ']; data.add_series('สัดส่วน',[R['sentiment'][k] for k in ('positive','neutral','negative')])
chart=s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Inches(.7),Inches(2.64), Inches(4.45), Inches(3.35),data).chart
chart.has_legend=False; chart.chart_style=10; chart.has_title=False; chart.plots[0].hole_size=72
for pt,c in zip(chart.series[0].points,[TEAL,'9BA9C0',RED]): pt.format.fill.solid(); pt.format.fill.fore_color.rgb=color(c); pt.format.line.fill.background()
text(s,'27',2.17,3.83,1.5,.72,40,True,INK,PP_ALIGN.CENTER); text(s,'ความคิดเห็น',1.9,4.56,2.0,.5,16,False,MUTED,PP_ALIGN.CENTER)
for i,(lbl,key,c) in enumerate([('บวก','positive',TEAL),('เป็นกลาง','neutral','9BA9C0'),('ลบ','negative',RED)]):
    x=.8+i*1.52; text(s,f"{R['sentiment'][key]:.1f}%",x,6.0,1.42,.42,24,True,c,PP_ALIGN.CENTER); text(s,lbl,x,6.45,1.42,.33,14,False,MUTED,PP_ALIGN.CENTER)
box(s,5.67,2.73,6.84,3.85,WHITE,radius=True)
text(s,'สิ่งที่ผลลัพธ์สะท้อน',5.97,2.99,6.24,.55,24,True)
text(s,'51.9% ถูกจัดเป็นความเห็นเชิงบวก\nหัวข้อที่พบมาก: นักกีฬา 55.6%\nรองลงมา: ผลงาน / ฟอร์มการเล่น 37.0%',5.97,3.8,6.14,1.68,23,False,INK)
text(s,'เป็นการสาธิตผลของโมเดลต่อข้อมูลชุดนี้\nไม่ใช่คะแนนความแม่นยำ หรือความเห็นของผู้ชมทั้งหมด',5.97,5.64,6.03,.76,16,True,RED)
note(s,'เวลาแนะนำ 60 วินาที อ้างอิง snapshot ของ video_id K01DQ6BSLHE ใน evidence.json ตัวเลข sentiment เท่ากับ 14/27, 4/27, 9/27 ตามผลจัดกลุ่มของโมเดล ประเด็นนักกีฬา 15/27 และผลงาน 10/27 ผู้จัดทำยังไม่ได้ตรวจ label กับผู้ประเมินมนุษย์ ห้ามพูดว่าแม่นยำ 51.9% และอย่าใช้จำนวน comment_count ปัจจุบันของ YouTube เป็นตัวหารแทน sample_size')

# 10 — testing
s=slide('ตรวจสอบการทำงานด้วยการทดสอบอัตโนมัติ','09 / SOFTWARE VERIFICATION')
box(s,.76,2.43,4.0,3.95,NAVY,radius=True)
text(s,'36 / 36',1.04,2.98,3.4,1.0,48,True,'9CE8D9')
text(s,'กรณีทดสอบผ่าน',1.07,4.12,3.38,.6,26,True,WHITE)
text(s,'pytest • ข้อมูล/API จำลอง\nฐานข้อมูลทดสอบแยกจากงานจริง',1.07,5.18,3.37,.85,16,False,'C6D2E8')
for y,t,b in [(2.54,'โครงสร้างและการคำนวณ','schema, comment ID, label และเปอร์เซ็นต์'),(3.85,'สถานการณ์ผิดพลาด','คีย์หาย, timeout, HTTP error และผลไม่ครบ'),(5.16,'การเชื่อมขั้นตอน','กรองวิดีโอกีฬาและบันทึกผลแบบจำลอง')]:
    label(s,'✓',5.3,y,.45,fill='D9F4EB',c=TEAL); text(s,t,6.0,y-.02,6.25,.55,22,True); text(s,b,6.0,y+.63,6.2,.62,17,False,MUTED)
note(s,'เวลาแนะนำ 50 วินาที ทดสอบจริงวันที่ 27 กันยายน 2569 exit code 0: 36 passed, 58 warnings in 1.19s รายงาน test-results.txt Warnings เป็น deprecation จึงยังมีงานดูแล dependency และเวลา 1.19 วินาทีเป็นเวลารัน tests ไม่ใช่ความเร็ววิเคราะห์จริง Tests จำลอง API จึงยืนยันการทำงานซอฟต์แวร์ ไม่ใช่ความแม่นยำ sentiment/emotion')

# 11 — deployment / latency
s=slide('เปิดใช้ผ่าน Cloudflare และเข้าใจเวลาที่ต้องรอ','10 / DEPLOYMENT & LATENCY')
for i,(title,body) in enumerate([('ผู้ใช้งาน','HTTPS / เบราว์เซอร์'),('Cloudflare Tunnel','เชื่อมมายังเครื่องผู้พัฒนา'),('เว็บ + Backend','เรียก YouTube และ Gemini')]):
    x=.75+i*4.2; box(s,x,2.43,3.45,1.38,WHITE,radius=True); text(s,title,x+.22,2.64,3.0,.46,22,True); text(s,body,x+.22,3.18,3.0,.38,15,False,MUTED)
    if i<2: arrow(s,x+3.61,2.95,.38,.27)
card(s,.77,4.35,5.78,2.25,'ปัจจัยด้านเวลา','ดึงคอมเมนต์ → รอโมเดล → บันทึกผล\n→ ดึงคอมเมนต์ยอดนิยมก่อนแสดงกราฟ',c=PURPLE)
card(s,6.82,4.35,5.75,2.25,'สภาพแวดล้อมสาธิต','Quick Tunnel เป็นลิงก์ชั่วคราว\nต้องเปิดเครื่องและเซิร์ฟเวอร์ตลอดเวลา',c=TEAL)
note(s,'เวลาแนะนำ 45 วินาที การใช้ Cloudflare ไม่ย้าย backend ไปโฮสต์บน Cloudflare ตัวแอปยังอยู่ที่เครื่องผู้พัฒนา การวิเคราะห์ขึ้นกับ API ภายนอก ปัจจุบันเปิดหน้าวิเคราะห์ใหม่จะเรียก Gemini อีกครั้งและกราฟยังรอคอมเมนต์ยอดนิยม ไม่มีผล benchmark p50/p95 จึงไม่ระบุตัวเลขความเร็ว')

# 12 — limits
s=slide('ข้อจำกัดที่ต้องใช้ประกอบการตีความผล','11 / LIMITATIONS')
for x,y,t,b,c in [(.76,2.4,'ขอบเขตกลุ่มตัวอย่าง','ไม่เกิน 100 ความคิดเห็นหลัก\nไม่รวมทุก reply และไม่ใช่การสุ่ม',PURPLE),(6.82,2.4,'ความหมายและบริบท','คำประชด คำย่อ และภาษากีฬา\nอาจทำให้จำแนกคลาดเคลื่อน',RED),(.76,4.58,'บริการภายนอก','ขึ้นกับคีย์ โควตา เครือข่าย\nและความพร้อมของ API',TEAL),(6.82,4.58,'การตรวจสอบย้อนหลัง','ยังไม่บันทึก prompt/model version\nและ label รายข้อความใน snapshot',PURPLE)]:
    card(s,x,y,5.74,1.98,t,b,c=c)
note(s,'เวลาแนะนำ 50 วินาที ย้ำว่ากราฟไม่ใช่ผลสำรวจประชากร และข้อความสรุปเป็นการสรุปความเห็น ไม่ยืนยันข้อกล่าวหาจากคอมเมนต์ ผลแต่ละครั้งอาจเปลี่ยนได้แม้ temperature ต่ำ ข้อมูลที่มีใน snapshot ยังไม่เพียงพอสำหรับทำซ้ำการทดลองอย่างสมบูรณ์')

# 13 — future evaluation
s=slide('แผนประเมินคุณภาพเชิงวิจัยในขั้นต่อไป','12 / EVALUATION PLAN','ส่วนนี้เป็นแผนดำเนินงาน ยังไม่มีค่าความแม่นยำหรือผลประเมินผู้ใช้')
card(s,.76,2.77,3.83,3.91,'สร้างข้อมูลอ้างอิง','เลือกตัวอย่างหลายประเภทกีฬา\nให้ผู้ประเมินอย่างน้อย 2 คน\nติดป้ายกำกับตามคู่มือเดียวกัน\nตรวจความสอดคล้องระหว่างคน','01',PURPLE)
card(s,4.78,2.77,3.83,3.91,'ประเมินผลจำแนก','เปรียบเทียบกับ label อ้างอิง\nรายงาน Macro-F1\nแสดง confusion matrix\nวิเคราะห์คำประชด / ภาษาปน','02',TEAL)
card(s,8.8,2.77,3.77,3.91,'ประเมินการใช้งาน','จับเวลาแยกแต่ละขั้นตอน\nรายงาน p50 / p95\nทดสอบภารกิจและความเข้าใจ\nของผู้ใช้กลุ่มเป้าหมาย','03',RED)
note(s,'เวลาแนะนำ 45 วินาที นี่คือข้อเสนอการประเมิน ไม่ใช่สิ่งที่ทำแล้ว ต้องกำหนดจำนวนตัวอย่าง แผนสุ่ม เกณฑ์การให้ label และข้อตกลงผู้ประเมินก่อนทดลอง การทดสอบ topics เป็น multi-label ต้องเลือก metric ที่เหมาะ เช่น micro/macro F1 และบันทึกเวอร์ชันโมเดลกับ prompt เพื่อทำซ้ำได้')

# 14 — conclusion
s=slide('สิ่งที่โครงงานพัฒนาได้ และแนวทางต่อยอด','13 / CONCLUSION',dark=True)
text(s,'จากคอมเมนต์ → ผลวิเคราะห์ที่อ่านและตรวจสอบได้',.78,2.36,11.8,.8,29,True,WHITE)
for x,n,t,b in [(.82,'01','เชื่อมข้อมูลจริง','YouTube + Gemini API'),(4.97,'02','ตรวจผลก่อนแสดง','schema + สถิติจากตัวอย่าง'),(9.12,'03','นำเสนอผ่านเว็บ','กราฟ + สรุปภาษาไทย')]:
    text(s,n,x,3.6,2.3,.75,42,True,'9CE8D9'); text(s,t,x,4.55,3.5,.63,23,True,WHITE); text(s,b,x,5.26,3.5,.65,17,False,'C6D2E8')
text(s,'ต่อยอด: ใช้ผลเดิมเมื่อเปิดซ้ำ • แสดงกราฟก่อนโหลดส่วนเสริม • ประเมินคุณภาพด้วยข้อมูลอ้างอิง',.85,6.32,11.72,.54,17,False,'B6C5DF')
note(s,'เวลาแนะนำ 40 วินาที สรุปผลงานเป็นระบบต้นแบบที่ใช้งานได้และมีการทดสอบการทำงาน อย่าอ้างความใหม่ด้านอัลกอริทึมหรือเหนือกว่าวิธีอื่น จุดต่อยอดเชิงวิศวกรรมคือใช้ snapshot เดิมและลดงานที่ต้องรอต่อเนื่อง จุดต่อยอดวิจัยคือประเมินกับข้อมูลอ้างอิง')

# 15 — references
s=slide('เอกสารอ้างอิงและหลักฐานประกอบ','REFERENCES')
refs=[('YouTube Data API — Search.list','https://developers.google.com/youtube/v3/docs/search/list'),('YouTube Data API — CommentThreads.list','https://developers.google.com/youtube/v3/docs/commentThreads/list'),('Gemini API — Structured outputs','https://ai.google.dev/gemini-api/docs/structured-output'),('Cloudflare — Quick Tunnels','https://developers.cloudflare.com/cloudflare-one/networks/connectors/cloudflare-tunnel/do-more-with-tunnels/trycloudflare/'),('AUCC — แนวทางและแม่แบบการนำเสนอ (ตรวจรุ่นงานก่อนส่ง)','https://aucc2026.rbru.ac.th/submission/submission-guideline/')]
for i,(title,url) in enumerate(refs):
    y=2.24+i*.73; text(s,f'[{i+1}]',.79,y,.6,.45,16,True,PURPLE)
    sh=text(s,title,1.42,y,10.8,.4,18,True); sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    text(s,url,1.42,y+.35,10.9,.34,10,False,MUTED)
text(s,'หลักฐานโครงงาน: source code • evidence.json • test-results.txt • ภาพหน้าจอระบบ\nวันที่ตรวจเอกสารและจัดทำภาพ: 27 กันยายน 2569',.81,6.16,11.7,.63,14,False,MUTED)
note(s,'แหล่งอ้างอิงเป็นเอกสารทางการ ลิงก์ชื่อรายการกดได้ หลักฐานท้องถิ่นอยู่ในโฟลเดอร์ presentation ภาพหน้าจอ UI ใช้ภาพปกและชื่อวิดีโอจาก YouTube เพื่อแสดงการทำงาน ไม่ใช่ผลงานภาพของผู้จัดทำ ไม่ได้อ้างว่าสไลด์ชุดนี้เป็นแม่แบบบังคับของ AUCC ต้องตรวจรูปแบบและกำหนดการของรุ่นงานที่จะสมัคร')

# 16 — discussion
s=slide('ขอบคุณครับ / ค่ะ','DISCUSSION',dark=True)
text(s,'คำถามและข้อเสนอแนะ',.78,2.48,11.85,1.0,42,True,WHITE)
text(s,'ระบบวิเคราะห์ความคิดเห็นต่อวิดีโอกีฬาบน YouTube ด้วย Gemini',.83,3.73,11.7,.75,24,False,'C6D2E8')
line(s,.85,4.91,5.6,4.91,TEAL,4)
text(s,'ประเด็นที่พร้อมอภิปราย\nขอบเขตตัวอย่าง • ความน่าเชื่อถือของผล • เวลาในการตอบสนอง',.85,5.51,11.62,1.0,20,False,'B6C5DF')
note(s,'เตรียมตอบคำถาม: 1) ทำไมไม่ train โมเดลเอง: ขอบเขตนี้เน้นประยุกต์ API และตรวจโครงสร้างผล 2) แม่นยำแค่ไหน: ยังไม่มี gold labels ไม่อ้าง accuracy 3) ทำไม 27 ทั้งที่หน้าวิดีโอมี 33: ใช้เฉพาะ top-level ที่ดึงได้ ณ เวลาวิเคราะห์ จำนวนหน้า YouTube อาจรวม replies/เปลี่ยนตามเวลา 4) ทำไม topics เกิน 100: หนึ่งคอมเมนต์มีหลายหัวข้อ 5) Cloudflare เก็บคีย์หรือไม่: คีย์อ่านจาก backend/.env และไม่ฝังใน frontend')

for s in prs.slides:
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.name=FONT
                rpr=r._r.get_or_add_rPr()
                rpr.set('lang','th-TH')
                for script in ('a:ea','a:cs'):
                    node=OxmlElement(script); node.set('typeface',FONT); rpr.append(node)

out=ROOT/'AUCC_Sports_Comment_Analysis.pptx'
prs.save(out)
speaker=['แนวทางนำเสนอประมาณ 10–12 นาที (เวลาสมมติสำหรับจัดเนื้อหา ไม่ใช่ข้อกำหนด AUCC)',
         'ก่อนส่ง: เติมชื่อผู้จัดทำ อาจารย์ที่ปรึกษา สาขา และมหาวิทยาลัยในสไลด์ 1',
         'ตรวจแม่แบบ ปีงาน และเวลานำเสนอของ AUCC รุ่นที่จะส่งอีกครั้ง',
         'สไลด์เป็นภาษาไทย ตัวอักษรและแผนภาพแก้ไขได้; ภาพ UI เป็นภาพหน้าจอจริง',
         'ห้ามใช้เปอร์เซ็นต์ผลจำแนกเป็นคะแนน accuracy และห้ามใช้เวลา pytest เป็น latency จริง','']
for i,n in notes: speaker += [f'สไลด์ {i:02d}',n,'']
(ROOT/'speaker-notes.txt').write_text('\n'.join(speaker),encoding='utf-8-sig')
print(f'Saved {out.name}: {len(prs.slides)} slides')
